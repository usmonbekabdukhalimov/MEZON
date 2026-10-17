"""Read-only audit. Requires openpyxl. Never modifies source workbooks."""
import argparse
from pathlib import Path
from collections import Counter
from decimal import Decimal
import hashlib
import json
import openpyxl

def audit(folder):
    results=[]
    all_ids=[]
    for name,header,id_col,amount_col,agent_col,client_col in [
        ('лапша.xlsx',1,0,5,2,4),('н.медов.xlsx',2,1,8,12,5),('крембер.xlsx',1,0,5,2,4)]:
        path=folder/name
        wb=openpyxl.load_workbook(path,data_only=False)
        ws=wb.worksheets[0]
        rows=[r for r in ws.iter_rows(min_row=header+1,values_only=True) if any(v is not None for v in r)]
        ids=[str(int(r[id_col])) if isinstance(r[id_col],(int,float)) and r[id_col]==int(r[id_col]) else str(r[id_col]) for r in rows]
        amounts=[Decimal(str(r[amount_col])) for r in rows]
        all_ids.extend(ids)
        results.append({'file':name,'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'sheets':wb.sheetnames,'sheet':ws.title,'range':ws.calculate_dimension(),'header_row':header,
            'headers':[c.value for c in ws[header]],'data_rows':len(rows),'unique_source_ids':len(set(ids)),
            'duplicate_id_rows':len(ids)-len(set(ids)),'summa_total':str(sum(amounts)),
            'nonpositive_amount_rows':sum(a<=0 for a in amounts),'fractional_amount_rows':sum(a!=int(a) for a in amounts),
            'distinct_agent_labels':len(set(r[agent_col] for r in rows)),
            'distinct_client_labels':len(set(r[client_col] for r in rows)),
            'repeated_agent_client_rows':sum(n-1 for n in Counter((r[agent_col],r[client_col]) for r in rows).values()),
            'missing_required_rows':sum(any(r[c] is None for c in [id_col,amount_col,agent_col,client_col]) for r in rows),
            'formula_cells':[c.coordinate for row in ws for c in row if c.data_type=='f'],
            'date_column_types':{ws.cell(header,c).value:dict(Counter(type(r[c-1]).__name__ for r in rows)) for c in ([4,5] if header==2 else [2])},
            'date_number_formats':{ws.cell(header,c).value:sorted(set(ws.cell(i,c).number_format for i in range(header+1,ws.max_row+1))) for c in ([4,5] if header==2 else [2])},
            'merged_ranges':[str(r) for r in ws.merged_cells.ranges]})
    return {'checked_on':'2026-10-18','method':'openpyxl 3.1.5 read-only inspection and Decimal sums; SHA-256 of original bytes',
        'files':results,'all_data_rows':sum(r['data_rows'] for r in results),'cross_file_repeated_id_count':len(all_ids)-len(set(all_ids)),
        'combined_summa_total':str(sum(Decimal(r['summa_total']) for r in results))}

if __name__=='__main__':
    p=argparse.ArgumentParser(); p.add_argument('folder',type=Path); p.add_argument('output',type=Path)
    a=p.parse_args(); result=audit(a.folder)
    a.output.parent.mkdir(parents=True,exist_ok=True)
    a.output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print('Audited',result['all_data_rows'],'rows. Total:',result['combined_summa_total'])
